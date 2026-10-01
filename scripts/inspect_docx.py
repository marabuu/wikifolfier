"""Inspect the .docx corpus: paragraph styles, headings, tables."""
import sys
from pathlib import Path

import docx

ROOT = Path(r"E:\dev\wikifolfier2\conteudo a adicionar")


def main() -> int:
    files = sorted(ROOT.rglob("*.docx"))
    print(f"{len(files)} arquivos .docx\n")

    for path in files[:6]:
        print("=" * 70)
        print(path.relative_to(ROOT))
        document = docx.Document(str(path))
        print(f"  paragrafos: {len(document.paragraphs)}  tabelas: {len(document.tables)}")

        styles = {}
        for paragraph in document.paragraphs:
            styles[paragraph.style.name] = styles.get(paragraph.style.name, 0) + 1
        print(f"  estilos: {dict(sorted(styles.items(), key=lambda kv: -kv[1])[:8])}")

        print("  --- primeiros blocos ---")
        shown = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            marker = "*" if paragraph.style.name.startswith("Heading") else " "
            print(f"   {marker} [{paragraph.style.name[:16]:<16}] {text[:78]}")
            shown += 1
            if shown >= 14:
                break
        if document.tables:
            table = document.tables[0]
            print(f"  --- 1a tabela: {len(table.rows)}x{len(table.columns)} ---")
            for row in table.rows[:3]:
                print("   | " + " | ".join(c.text.strip()[:22] for c in row.cells))

    return 0


if __name__ == "__main__":
    sys.exit(main())