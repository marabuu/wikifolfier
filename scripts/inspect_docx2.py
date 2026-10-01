"""Look at the multi-item homebrew files to find their internal structure."""
import sys
from pathlib import Path

import docx

ROOT = Path(r"E:\dev\wikifolfier2\conteudo a adicionar")
TARGETS = [
    "Regras Homebrew/Novos feitiços.docx",
    "Regras Homebrew/Novos Feats.docx",
    "Regras Homebrew/Novos fighting styles.docx",
    "Regras Homebrew/Regras Gerais.docx",
    "Regras Homebrew/Bastião.docx",
]


def main() -> int:
    for relative in TARGETS:
        path = ROOT / relative
        if not path.exists():
            print(f"!! ausente: {relative}")
            continue
        print("=" * 72)
        print(relative)
        document = docx.Document(str(path))
        print(f"  paragrafos={len(document.paragraphs)} tabelas={len(document.tables)}")
        shown = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            runs_bold = any(run.bold for run in paragraph.runs if run.bold)
            flag = "BOLD" if runs_bold else "    "
            print(f"   {flag} | {text[:88]}")
            shown += 1
            if shown >= 26:
                break
        if document.tables:
            table = document.tables[0]
            print(f"   [tabela {len(table.rows)}x{len(table.columns)}]")
            for row in table.rows[:3]:
                print("     | " + " | ".join(c.text.strip()[:18] for c in row.cells))
    return 0


if __name__ == "__main__":
    sys.exit(main())